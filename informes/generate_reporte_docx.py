"""
Generador de Informe Técnico Académico en Formato DOCX (Word)
Norma Internacional ISO/IEC/IEEE 29119 - Automatización de Pruebas E2E con Playwright
Grupo 02 - Calidad de Software UNSA 2026
"""

import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Aplica color de fondo hexadecimal a una celda de tabla."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Establece márgenes internos (padding) de celda en dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    """Aplica bordes sutiles y profesionales a la tabla."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_report():
    doc = Document()

    # Configuración de márgenes estándar de 1 pulgada (Letter)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Colores Institucionales
    NAVY = RGBColor(15, 23, 42)      # #0f172a
    BLUE = RGBColor(30, 58, 138)     # #1e3a8a
    MUTED = RGBColor(71, 85, 105)    # #475569
    DARK = RGBColor(30, 41, 59)      # #1e293b

    # =========================================================================
    # 1. PORTADA INSTITUCIONAL FORMAL
    # =========================================================================
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("UNIVERSIDAD NACIONAL DE SAN AGUSTÍN DE AREQUIPA")
    r_uni.font.name = "Arial"
    r_uni.font.size = Pt(14)
    r_uni.font.bold = True
    r_uni.font.color.rgb = NAVY

    p_fac = doc.add_paragraph()
    p_fac.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fac.paragraph_format.space_after = Pt(24)
    r_fac = p_fac.add_run("FACULTAD DE INGENIERÍA DE PRODUCCIÓN Y SERVICIOS\nESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS")
    r_fac.font.name = "Arial"
    r_fac.font.size = Pt(11)
    r_fac.font.bold = True
    r_fac.font.color.rgb = MUTED

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("INFORME TÉCNICO DE AUTOMATIZACIÓN DE PRUEBAS DE SOFTWARE E2E BAJO LA NORMA INTERNACIONAL ISO/IEC/IEEE 29119")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = BLUE

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(40)
    r_sub = p_sub.add_run("Evaluación Comparativa de Calidad: Sistema de Vigilancia Epidemiológica IRA Perú (ISPySA-Pneumonia) y Plataforma E-Commerce SauceDemo")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = MUTED

    # Cuadro de Identificación del Grupo y Curso
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    data_info = [
        ("Asignatura:", "Calidad de Software (Semestre 2026-B)"),
        ("Trabajo Parcial:", "Trabajo Parcial 6 (TP6) - Automatización de Pruebas ISO 29119"),
        ("Equipo de Trabajo:", "Grupo 02"),
        ("Integrantes:", "• Condori Ayna, Yoel Herbert\n• Márquez Choquepuma, Marco Antonio\n• Pérez Luque, Jeremy Frank\n• Nina Cusirramos, Rafael Alexander\n• Arce Manrique, Leonardo Gabriel"),
        ("Lugar y Fecha:", "Arequipa, Octubre de 2026")
    ]

    col_widths = [Inches(2.2), Inches(4.3)]
    for i, (label, val) in enumerate(data_info):
        row = info_table.rows[i]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = col_widths[0], col_widths[1]

        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = NAVY

        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK

        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, top=100, bottom=100, left=120, right=120)
        set_cell_margins(cell_val, top=100, bottom=100, left=120, right=120)

    set_table_borders(info_table, "CBD5E1")

    doc.add_page_break()

    # =========================================================================
    # 2. RESUMEN EJECUTIVO Y CONTROL DE VERSIONES
    # =========================================================================
    h1 = doc.add_heading("RESUMEN EJECUTIVO", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(8)

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(12)
    p_exec.add_run(
        "El presente informe documenta el diseño, implementación, ejecución y evaluación de una suite de pruebas "
        "automatizadas End-to-End (E2E) bajo los estándares formales de la familia de normas ISO/IEC/IEEE 29119 "
        "(Partes 1 a 5). Se implementó una arquitectura de verificación basada en Playwright y Chromium sobre dos sistemas "
        "de software complementarios: (a) el Sistema de Información para la Toma de Decisiones en Salud Pública "
        "(ISPySA-Pneumonia, Perú 2000-2023), que incorpora el Canal Endémico de Bortman (OPS/MINSA), el Monitor de Triaje "
        "Manchester y la Matriz Bioclimática Regional de Heladas y Friaje; y (b) la plataforma de comercio electrónico "
        "SauceDemo, utilizada como marco referencial de contraste funcional. Se diseñaron 24 casos de prueba formalmente "
        "trazados a 6 requisitos de ingeniería (R1 a R6). La ejecución rigurosa de 3 corridas consecutivas evidenció un "
        "100% de tasa de aprobación (24/24 pruebas aprobadas por corrida), con una tasa de flakiness de 0.00% y un "
        "análisis de retorno de inversión (ROI) que demuestra la recuperación del costo en 7 ciclos de regresión."
    )

    # Tabla de Control de Versiones
    doc.add_heading("Control de Versiones y Metadatos del Documento", level=2)
    tbl_ver = doc.add_table(rows=3, cols=5)
    tbl_ver.alignment = WD_TABLE_ALIGNMENT.CENTER
    ver_headers = ["Versión", "Fecha", "Autor(es)", "Revisión", "Estado"]
    ver_rows = [
        ver_headers,
        ["1.0.0", "18/10/2026", "Grupo 02 (Condori, Márquez, Pérez, Nina, Arce)", "Ingeniería de Calidad", "Aprobado"],
        ["1.1.0", "19/10/2026", "Grupo 02 (Condori, Márquez, Pérez, Nina, Arce)", "Aseguramiento ISO 29119", "Oficial Definitivo"]
    ]
    for row_idx, r_data in enumerate(ver_rows):
        row = tbl_ver.rows[row_idx]
        for col_idx, val in enumerate(r_data):
            c = row.cells[col_idx]
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if row_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(c, "1E3A8A")
            else:
                set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    set_table_borders(tbl_ver, "CBD5E1")

    # =========================================================================
    # 3. SECCIÓN 1: MARCO TEÓRICO Y CONCEPTUAL (ISO/IEC/IEEE 29119)
    # =========================================================================
    h1 = doc.add_heading("1. MARCO NORMATIVO INTERNACIONAL: SERIE ISO/IEC/IEEE 29119", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)

    p_norm = doc.add_paragraph()
    p_norm.paragraph_format.line_spacing = 1.15
    p_norm.add_run(
        "La norma internacional ISO/IEC/IEEE 29119 reemplaza los estándares históricos IEEE 829 (documentación) e "
        "IEEE 1008 (pruebas unitarias), unificando el ciclo de vida de pruebas de software en un cuerpo normativo "
        "coherente estructurado en cinco partes fundamentales:"
    )

    norm_points = [
        ("ISO/IEC/IEEE 29119-1 (Conceptos y Definiciones): ",
         "Establece el vocabulario estandarizado, el modelo ontológico de pruebas de software y los principios rectores "
         "que rigen la verificación formal, definiendo conceptos clave como falla, defecto, error y cobertura de prueba."),
        ("ISO/IEC/IEEE 29119-2 (Procesos de Prueba): ",
         "Define un modelo de procesos en tres capas: (a) Gobernanza y políticas organizacionales de prueba; "
         "(b) Procesos de gestión del proyecto de prueba (Planificación, Monitoreo y Control, y Cierre); y "
         "(c) Procesos dinámicos de prueba (Diseño, Implementación, Configuración de Entorno, Ejecución y Reporte de Incidentes)."),
        ("ISO/IEC/IEEE 29119-3 (Documentación de Pruebas): ",
         "Normaliza la estructura obligatoria de los artefactos de prueba, incluyendo el Plan Organizacional de Pruebas, "
         "el Plan del Proyecto de Pruebas (Test Plan), la Especificación de Diseño de Pruebas (Test Design Specification), "
         "la Especificación de Casos de Prueba (Test Case Specification), el Procedimiento de Ejecución (Test Procedure) y "
         "el Reporte de Estado y Cierre de Pruebas (Test Summary Report)."),
        ("ISO/IEC/IEEE 29119-4 (Técnicas de Diseño de Pruebas): ",
         "Formaliza las técnicas analíticas basadas en especificación (caja negra), basadas en estructura (caja blanca) y "
         "basadas en experiencia. En este proyecto se aplican particionamiento de equivalencia, análisis de valores límite (BVA) "
         "y validación de transiciones de estado."),
        ("ISO/IEC/IEEE 29119-5 (Pruebas Guiadas por Palabras Clave): ",
         "Define la arquitectura modular para frameworks Keyword-Driven y Data-Driven, permitiendo la desacoplación "
         "entre la lógica de automatización y las especificaciones funcionales de negocio.")
    ]

    for title, desc in norm_points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        rt = p.add_run(title)
        rt.font.bold = True
        rt.font.color.rgb = BLUE
        rd = p.add_run(desc)
        rd.font.size = Pt(10)

    # =========================================================================
    # 4. SECCIÓN 2: ALCANCE DE PRUEBAS Y ESPECIFICACIÓN DE REQUISITOS (R1 A R6)
    # =========================================================================
    h1 = doc.add_heading("2. ESPECIFICACIÓN DE REQUISITOS Y ALCANCE DE PRUEBAS (R1 A R6)", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)

    p_req = doc.add_paragraph()
    p_req.paragraph_format.line_spacing = 1.15
    p_req.add_run(
        "Conforme a las directrices de la práctica y la consigna del Trabajo Parcial 6, la evaluación de calidad se "
        "extendió de forma dual sobre dos sistemas de software:"
    )

    # Tabla de Requisitos R1 a R6
    tbl_req = doc.add_table(rows=7, cols=3)
    tbl_req.alignment = WD_TABLE_ALIGNMENT.CENTER
    req_data = [
        ["Requisito", "Sistema A: Salud Pública (ISPySA-Pneumonia)", "Sistema B: E-Commerce (SauceDemo)"],
        ["R1: Control de Acceso y Contexto", "Segmentación demográfica pediátrica (men5) y geriátrica (60mas); actualización contextual en hero banner.", "Autenticación segura con credenciales estándar y detección de usuario bloqueado (locked_out_user)."],
        ["R2: Alerta Sanitaria / Inspección", "Semáforo epidemiológico dinámico y Corredor Endémico de Bortman (OPS/MINSA) en 4 zonas.", "Inspección de integridad estructural del catálogo (6 productos con títulos y precios positivos)."],
        ["R3: Demanda Hospitalaria / Ordenamiento", "Monitor de Triaje Manchester y proyección de camas UCI y consumo diario de oxígeno medicinal.", "Clasificación de productos por precio ascendente (lohi) y descendente (hilo) verificando P[i] <= P[i+1]."],
        ["R4: Estratificación / Carrito", "Estratificación de persistencia en Top 3 y Matriz Bioclimática Regional (Heladas y Friaje).", "Gestión reactiva de carrito de compras: adición de ítems, actualización de badge y supresión parcial."],
        ["R5: Consistencia Aritmética", "Validación matemática defensiva de tasas HR y CFR con tolerancia absoluta <= 0.05%.", "Validación obligatoria de formulario de despacho y transición a resumen de checkout."],
        ["R6: Modelos Predictivos / Cierre", "Evaluación de modelos ML (XGBoost/RF) y matriz de gobernanza ISO 9241 e ISO/IEC 25010.", "Cálculo aritmético de subtotal, impuesto (8%) y total neto; finalización de orden y logout seguro."]
    ]

    for row_idx, r_data in enumerate(req_data):
        row = tbl_req.rows[row_idx]
        for col_idx, val in enumerate(r_data):
            c = row.cells[col_idx]
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if row_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(c, "1E3A8A")
            else:
                set_cell_background(c, "F1F5F9" if col_idx == 0 else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF"))
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    set_table_borders(tbl_req, "CBD5E1")

    # =========================================================================
    # 5. SECCIÓN 3: MATRIZ DE CASOS DE PRUEBA Y TRAZABILIDAD (T01 A T12)
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("3. MATRIZ DE DISEÑO DE CASOS DE PRUEBA Y TRAZABILIDAD (ISO 29119-3/4)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(8)

    doc.add_heading("3.1 Suite de Casos de Prueba - Sistema A: ISPySA-Pneumonia (Salud Pública)", level=2)
    tbl_t_ispysa = doc.add_table(rows=13, cols=5)
    tbl_t_ispysa.alignment = WD_TABLE_ALIGNMENT.CENTER
    ispysa_cases = [
        ["ID", "Req.", "Técnica ISO 29119-4", "Descripción del Caso de Prueba", "Criterio de Aceptación"],
        ["T01", "R1", "Partición de Equivalencia", "Carga inicial y verificación de contexto men5", "Banner contextual y 4 tarjetas KPI visibles."],
        ["T02", "R1", "Transición de Estados", "Conmutación interactiva a cohorte 60mas", "Recálculo dinámico de indicadores y tarjeta CFR."],
        ["T03", "R2", "Partición de Equivalencia", "Evaluación de alerta sanitaria canalizada", "Banner Bortman categoriza según umbral Q1, Q2 o Q3."],
        ["T04", "R2", "Análisis Valores Límite", "Renderizado del Corredor de Bortman en Tab I", "Visualización de 4 zonas y 4 tarjetas de cuartiles."],
        ["T05", "R3", "Inspección Funcional", "Curvas de severidad clínica en Tab II", "Gráfico dual HR/CFR y dispersión cuadrante visibles."],
        ["T06", "R3", "Modelado Estadístico", "Monitor de Triaje Manchester y recursos críticos", "Demanda UCI > 0, consumo O2 > 0 y 5 prioridades MTS."],
        ["T07", "R4", "Partición de Equivalencia", "Persistencia departamental Top 3 en Tab III", "Gráfico de barras de persistencia histórica desplegado."],
        ["T08", "R4", "Análisis de Dominio", "Matriz Bioclimática Regional (Heladas/Friaje)", "Clasificación de 25 departamentos por factor térmico."],
        ["T09", "R5", "Consistencia Aritmética", "Verificación matemática de tasa HR", "|HR_reportado - (hosp/casos)*100| <= 0.05%."],
        ["T10", "R5", "Consistencia Aritmética", "Verificación matemática de tasa CFR", "|CFR_reportado - (death/casos)*100| <= 0.05%."],
        ["T11", "R6", "Inspección de Dominio", "Métricas de precisión ML en Tab IV", "MAE, RMSE, R2 y modelo óptimo XGBoost visibles."],
        ["T12", "R6", "Auditoría Normativa", "Matriz de cumplimiento ISO 9241 en Tab V", "Presencia explícita de normas 11, 210, 110, 12 y 25010."]
    ]

    for row_idx, r_data in enumerate(ispysa_cases):
        row = tbl_t_ispysa.rows[row_idx]
        for col_idx, val in enumerate(r_data):
            c = row.cells[col_idx]
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if row_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(c, "1E3A8A")
            else:
                set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
    set_table_borders(tbl_t_ispysa, "CBD5E1")

    doc.add_heading("3.2 Suite de Casos de Prueba - Sistema B: SauceDemo (E-Commerce)", level=2)
    tbl_t_sauce = doc.add_table(rows=13, cols=5)
    tbl_t_sauce.alignment = WD_TABLE_ALIGNMENT.CENTER
    sauce_cases = [
        ["ID", "Req.", "Técnica ISO 29119-4", "Descripción del Caso de Prueba", "Criterio de Aceptación"],
        ["T01", "R1", "Partición de Equivalencia", "Login exitoso con usuario válido standard_user", "Redirección a /inventory.html y título 'Products'."],
        ["T02", "R1", "Manejo de Errores", "Bloqueo de acceso para locked_out_user", "Mensaje 'Sorry, this user has been locked out.' visible."],
        ["T03", "R2", "Inspección Estructural", "Integridad del catálogo comercial", "Exactamente 6 productos con nombre y precio > 0."],
        ["T04", "R2", "Transición de Estados", "Navegación al detalle y retorno al inventario", "Detalle de 'Sauce Labs Backpack' ($29.99) y retorno fluido."],
        ["T05", "R3", "Análisis de Valores Límite", "Ordenamiento ascendente de precios (lohi)", "Verificación matemática: P[i] <= P[i+1] para todo i."],
        ["T06", "R3", "Análisis de Valores Límite", "Ordenamiento descendente de precios (hilo)", "Verificación matemática: P[i] >= P[i+1] para todo i."],
        ["T07", "R4", "Transición de Estados", "Adición de ítem al carrito de compras", "Badge = '1' y botón conmuta a 'Remove'."],
        ["T08", "R4", "Transición de Estados", "Adición múltiple y supresión parcial de ítems", "Badge transita de '2' a '1' reactivamente."],
        ["T09", "R5", "Prueba Negativa / BVA", "Validación de campos vacíos en Checkout Step One", "Mensaje 'Error: First Name is required' desplegado."],
        ["T10", "R5", "Partición de Equivalencia", "Ingreso formal de datos válidos de despacho", "Transición exitosa a /checkout-step-two.html."],
        ["T11", "R6", "Consistencia Aritmética", "Verificación financiera: Subtotal + Tax = Total", "Subtotal exacto, Tax = Subtotal*0.08, Total consistente."],
        ["T12", "R6", "Transición de Estados", "Finalización de orden y logout seguro", "Mensaje 'Thank you for your order!' y logout limpio."]
    ]

    for row_idx, r_data in enumerate(sauce_cases):
        row = tbl_t_sauce.rows[row_idx]
        for col_idx, val in enumerate(r_data):
            c = row.cells[col_idx]
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if row_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(c, "1E3A8A")
            else:
                set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
    set_table_borders(tbl_t_sauce, "CBD5E1")

    # =========================================================================
    # 6. SECCIÓN 4: PROTOCOLO DE EJECUCIÓN Y ANÁLISIS DE REPRODUCIBILIDAD (3 CORRIDAS)
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("4. PROTOCOLO DE EJECUCIÓN Y ANÁLISIS DE REPRODUCIBILIDAD (3 CORRIDAS)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(8)

    p_proto = doc.add_paragraph()
    p_proto.paragraph_format.line_spacing = 1.15
    p_proto.add_run(
        "Para verificar la estabilidad temporal, ausencia de condiciones de carrera y reproducibilidad determinista "
        "del framework automatizado (ISO/IEC/IEEE 29119-2, Cláusula de Control de Ejecución), se ejecutó el protocolo "
        "de tres corridas secuenciales completas en Chromium headless bajo parámetros estandarizados (workers: 1, retries: 0)."
    )

    tbl_runs = doc.add_table(rows=5, cols=7)
    tbl_runs.alignment = WD_TABLE_ALIGNMENT.CENTER
    runs_headers = ["Corrida", "Fecha y Hora", "Total Pruebas", "Aprobadas", "Fallidas", "Duración", "Flakiness"]
    runs_data = [
        runs_headers,
        ["Corrida 1", "19/10/2026 21:15", "24", "24 (100%)", "0 (0%)", "69.51 s", "0.00%"],
        ["Corrida 2", "19/10/2026 21:16", "24", "24 (100%)", "0 (0%)", "68.84 s", "0.00%"],
        ["Corrida 3", "19/10/2026 21:17", "24", "24 (100%)", "0 (0%)", "68.20 s", "0.00%"],
        ["Consolidado", "Promedio Global", "72", "72 (100%)", "0 (0%)", "68.85 s", "0.00%"]
    ]

    for row_idx, r_data in enumerate(runs_data):
        row = tbl_runs.rows[row_idx]
        for col_idx, val in enumerate(r_data):
            c = row.cells[col_idx]
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if row_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(c, "1E3A8A")
            elif row_idx == 4:
                r.font.bold = True
                set_cell_background(c, "E2E8F0")
            else:
                set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=70, bottom=70, left=90, right=90)
    set_table_borders(tbl_runs, "CBD5E1")

    p_flaky = doc.add_paragraph()
    p_flaky.paragraph_format.space_before = Pt(8)
    p_flaky.paragraph_format.line_spacing = 1.15
    p_flaky.add_run(
        "Análisis de Flakiness Formal:\n"
        "La tasa de inconsistencia inter-ejecución (Flakiness Rate) se calcula como:\n"
        "Flakiness = (Casos con Resultado Discrepante entre Corridas / Total de Pruebas Evaluadas) × 100% = (0 / 24) = 0.00%.\n"
        "Esto certifica la robustez de los selectores semánticos y la estabilidad del pipeline de aseguramiento de calidad."
    )

    # =========================================================================
    # 7. SECCIÓN 5: ANÁLISIS ECONÓMICO Y RETORNO DE INVERSIÓN (ROI)
    # =========================================================================
    h1 = doc.add_heading("5. ANÁLISIS ECONÓMICO Y RETORNO DE INVERSIÓN (ROI DE AUTOMATIZACIÓN)", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)

    p_roi = doc.add_paragraph()
    p_roi.paragraph_format.line_spacing = 1.15
    p_roi.add_run(
        "Conforme al estándar ISO/IEC/IEEE 29119-2 (Gestión Organizacional y Viabilidad de Pruebas), la adopción de "
        "un framework de pruebas automatizadas requiere una justificación cuantitativa del costo versus beneficio:\n\n"
        "Parámetros del Modelo Financiero:\n"
        "• Costo de Diseño e Implementación Inicial (C_setup): 180 minutos (3 horas hombre de ingeniería).\n"
        "• Costo de Ejecución Manual por Ciclo de Regresión (C_manual): 28 minutos (14 min por sistema × 2 sistemas).\n"
        "• Costo de Ejecución Automatizada por Ciclo (C_auto): 2 minutos (supervisión y revisión de reporte HTML).\n"
        "• Ahorro Neto por Ciclo de Regresión (ΔC): C_manual - C_auto = 28 min - 2 min = 26 minutos por ciclo.\n\n"
        "Ecuación del Punto de Equilibrio (Break-Even Point):\n"
        "N_breakeven = ceil( C_setup / ΔC ) = ceil( 180 / 26 ) = ceil( 6.92 ) = 7 ciclos de regresión.\n\n"
        "A partir del séptimo ciclo de ejecución en el pipeline de CI/CD, la inversión inicial de automatización queda "
        "100% amortizada. En una proyección conservadora de 50 ciclos de despliegue, el ahorro neto acumulado asciende a:\n"
        "Ahorro_50 = (50 × 26 min) - 180 min = 1,300 min - 180 min = 1,120 minutos (18.67 horas hombre liberadas)."
    )

    tbl_roi = doc.add_table(rows=5, cols=5)
    tbl_roi.alignment = WD_TABLE_ALIGNMENT.CENTER
    roi_headers = ["Ciclos de Regresión (N)", "Tiempo Manual (min)", "Tiempo Automatizado (min)", "Costo Acumulado Total (min)", "Ahorro Neto Acumulado (min)"]
    roi_data = [
        roi_headers,
        ["1 ciclo", "28 min", "182 min", "182 min", "-154 min (Inversión inicial)"],
        ["7 ciclos (Punto de Equilibrio)", "196 min", "194 min", "194 min", "+2 min (Amortización completada)"],
        ["20 ciclos", "560 min", "220 min", "220 min", "+340 min (5.67 horas hombre)"],
        ["50 ciclos", "1,400 min", "280 min", "280 min", "+1,120 min (18.67 horas hombre)"]
    ]

    for row_idx, r_data in enumerate(roi_data):
        row = tbl_roi.rows[row_idx]
        for col_idx, val in enumerate(r_data):
            c = row.cells[col_idx]
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if row_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(c, "1E3A8A")
            elif row_idx == 2:
                r.font.bold = True
                set_cell_background(c, "FEF3C7")
            else:
                set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=70, bottom=70, left=90, right=90)
    set_table_borders(tbl_roi, "CBD5E1")

    # =========================================================================
    # 8. SECCIÓN 6: DICTAMEN TÉCNICO DE ACEPTACIÓN Y CONCLUSIONES ISO 29119
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("6. DICTAMEN TÉCNICO DE ACEPTACIÓN Y CONCLUSIONES (ISO 29119-2/3)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(8)

    p_dict = doc.add_paragraph()
    p_dict.paragraph_format.line_spacing = 1.15
    p_dict.add_run(
        "Evaluación de Criterios de Aceptación (Exit Criteria):\n"
        "1. Cobertura de Requisitos: 100% de los requisitos estipulados (R1 a R6) verificados en ambos sistemas.\n"
        "2. Tasa de Éxito de Pruebas: 100% de casos de prueba aprobados (24/24 en cada una de las 3 corridas).\n"
        "3. Tasa de Flakiness: 0.00% de inconsistencias inter-ejecución.\n"
        "4. Defectos Bloqueantes Residuales: Cero (0) defectos abiertos de severidad crítica o alta.\n\n"
        "DICTAMEN FORMAL DE CALIDAD:\n"
        "El equipo evaluador del Grupo 02 emite dictamen de APROBACIÓN TÉCNICA SATISFACTORIA sin reservas para el "
        "pase a producción del Sistema de Vigilancia Epidemiológica IRA Perú (ISPySA-Pneumonia) y certifica la plena "
        "conformidad del proceso de verificación bajo la norma internacional ISO/IEC/IEEE 29119.\n\n"
        "Conclusiones Técnicas Principales:\n"
        "• La modernización de ISPySA-Pneumonia con el Canal Endémico de Bortman (OPS/MINSA), el Triaje Manchester "
        "y la Matriz Bioclimática elevó la pertinencia clínica del aplicativo a estándares de vigilancia epidemiológica internacional.\n"
        "• La implementación de selectores semánticos estables erradicó el acoplamiento frágil a la interfaz visual, "
        "permitiendo tiempos de ejecución inferiores a 70 segundos para la suite completa de 24 casos de prueba.\n"
        "• La amortización de la inversión en solo 7 ciclos confirma que la automatización de pruebas E2E con Playwright "
        "constituye la estrategia óptima para garantizar la calidad continua del software institucional."
    )

    output_path = Path(__file__).resolve().parents[1] / "informes" / "trabajo_6_iso29119_reporte.docx"
    doc.save(str(output_path))
    print(f"Informe formal DOCX generado exitosamente en: {output_path}")

if __name__ == "__main__":
    build_report()
