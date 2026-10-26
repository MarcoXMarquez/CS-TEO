import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    """Aplica color de fondo a una celda de tabla en Word."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Configura márgenes internos de una celda (en dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_placeholder(doc, fig_num, title, instructions):
    """Crea una caja de placeholder formateada profesionalmente para capturas de pantalla."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC") # Slate 50
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border styling (Left thick border blue, others light)
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        r'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        r'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run_tag = p.add_run(f"📷 [FIGURA {fig_num}: PLACEHOLDER DE CAPTURA]\n")
    run_tag.bold = True
    run_tag.font.name = "Arial"
    run_tag.font.size = Pt(9.5)
    run_tag.font.color.rgb = RGBColor(37, 99, 235) # Blue 600
    
    run_title = p.add_run(f"{title}\n\n")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    
    run_desc = p.add_run(f"Instrucciones para la captura:\n{instructions}")
    run_desc.italic = True
    run_desc.font.name = "Arial"
    run_desc.font.size = Pt(8.5)
    run_desc.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
    
    # Espaciador después de la tabla
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(4)

def format_paragraph(p, space_before=0, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_ieee(doc, text, level=1):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=12 if level==1 else 8, space_after=4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.bold = True
    if level == 1:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    elif level == 2:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(30, 64, 175) # Blue 800
    else:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    return p

print("Helper functions compiled.")
